"""Markdown (article assembler output) → .docx with python-docx: headings, paragraphs with
inline **bold**/*italic*/`code`, bullet lists, images with captions, pipe tables (wide ones on
landscape pages), reference list. A4, serif. Only the constructs used by the manuscript are handled."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from tools.paper3_audit.article import ARTICLE_DIR

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
LANDSCAPE_FROM_COLS = 9


def add_inline(par, text: str, size: Pt | None = None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
        else:
            r = par.add_run(tok)
        if size is not None and r.font.size is None:
            r.font.size = size


def _set_orientation(section, landscape: bool):
    w, h = Cm(21.0), Cm(29.7)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (h, w) if landscape else (w, h)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))


def add_pipe_table(doc, lines: list[str]) -> int:
    header = [c.strip() for c in lines[0].strip().strip("|").split("|")]
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines[2:]]
    n = len(header)
    size = Pt(9) if n <= 6 else Pt(7.5) if n <= 10 else Pt(6.5)
    landscape = n >= LANDSCAPE_FROM_COLS
    if landscape:
        _set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), True)
    table = doc.add_table(rows=1, cols=n)
    table.style = "Table Grid"
    table.autofit = True
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h); r.bold = True; r.font.size = size
    for row in rows:
        cells = table.add_row().cells
        for i in range(n):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(row[i] if i < len(row) else "")
            r.font.size = size
    doc.add_paragraph()
    if landscape:
        _set_orientation(doc.add_section(WD_SECTION.NEW_PAGE), False)
    return len(rows)


def build(md_path: Path, out_path: Path) -> dict:
    doc = Document()
    _set_orientation(doc.sections[0], False)
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i, n = 0, len(lines)
    para_buf: list[str] = []
    stats = {"tables": 0, "table_rows": 0}

    def flush():
        nonlocal para_buf
        if para_buf:
            text = " ".join(l.strip() for l in para_buf)
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            add_inline(p, text)
            para_buf = []

    while i < n:
        line = lines[i]
        if line.startswith("<!--"):
            while i < n and "-->" not in lines[i]:
                i += 1
            i += 1; continue
        if not line.strip():
            flush(); i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            flush(); level = len(m.group(1))
            h = doc.add_heading(level=0 if level == 1 else level - 1); add_inline(h, m.group(2))
            i += 1; continue
        m = re.match(r"^!\[(Fig(?:S)?\d{2})\]\((figures/[^)]+)\)", line)
        if m:
            flush()
            img = md_path.parent / m.group(2)
            if img.exists():
                doc.add_picture(str(img), width=Cm(16))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1; continue
        if line.startswith("|"):
            flush()
            block = []
            while i < n and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            if len(block) >= 2:
                stats["tables"] += 1
                stats["table_rows"] += add_pipe_table(doc, block)
            continue
        if line.startswith("- "):
            flush()
            while i < n and lines[i].startswith("- "):
                item = lines[i][2:]
                i += 1
                while i < n and lines[i].startswith("  ") and lines[i].strip():
                    item += " " + lines[i].strip(); i += 1
                p = doc.add_paragraph(style="List Bullet"); add_inline(p, item, Pt(10))
            continue
        if re.match(r"^\*\*(Fig(?:S)?\d{2}|T\d{2}[a-c]?)\b", line) and para_buf == []:
            cap = [line]; i += 1
            while i < n and lines[i].strip() and not lines[i].startswith("|"):
                cap.append(lines[i]); i += 1
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8)
            add_inline(p, " ".join(c.strip() for c in cap), Pt(9.5))
            continue
        para_buf.append(line); i += 1
    flush()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return stats


def run(stem: str = "Paper3_final") -> int:
    md = ARTICLE_DIR / f"{stem}.md"
    out = ARTICLE_DIR / f"{stem}.docx"
    stats = build(md, out)
    d = Document(str(out))
    n_img = sum(1 for _ in d.inline_shapes)
    print(f"docx: {out.name} — {len(d.paragraphs)} paragraphs, {n_img} images, {len(d.tables)} tables "
          f"({stats['table_rows']} rows), {len(d.sections)} sections, {out.stat().st_size/1e6:.1f} MB")
    return 0
