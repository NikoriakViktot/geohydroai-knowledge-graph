"""assemble: build the manuscript from the paper repository's templates, own numbers and tables.

Dialects (manifest ``placeholders.dialect``):
- ``ghai``            section templates + OWN_EVIDENCE claims + citations (src/workbench/assemble/render.py);
- ``floodstate_fill`` one template whose numbers are table cells (src/workbench/assemble/fill.py);
- ``none``            the manuscript is written directly; nothing to assemble.

Staged for delivery: the manuscript (``paths.manuscript_out``, English), OPEN_ITEMS.md and
assembly_report.json beside it, claims_used.csv (ghai), and a .docx in an ignored private/ folder
(or ``assembly.docx``). ``--final`` is the submission gate: with any [PENDING] marker or
[[MISSING]] cell left, nothing is staged and the step fails.
"""

from __future__ import annotations

import io
import json
import tempfile
from pathlib import Path, PurePosixPath

import pandas as pd
import yaml

from src.workbench.assemble import fill as F
from src.workbench.assemble import markers
from src.workbench.assemble.render import AssemblyError, Assembler, Inputs, open_items, parse_captions
from src.workbench.delivery import sha256
from src.workbench.steps import Context, provenance, stage

IMAGE_PREFERENCE = (".png", ".jpg", ".jpeg", ".svg", ".pdf")


def _under(files: dict[str, bytes], directory: str) -> dict[str, bytes]:
    d = directory.rstrip("/") + "/"
    return {p: b for p, b in files.items() if p.startswith(d)}


def _templates(files: dict[str, bytes], path: str) -> list[tuple[str, str]]:
    if path in files:
        return [(PurePosixPath(path).name, files[path].decode("utf-8"))]
    found = sorted((p, b) for p, b in _under(files, path).items()
                   if p.endswith(".md") and not PurePosixPath(p).name.startswith("_") and p.count("/") == path.rstrip("/").count("/") + 1)
    return [(PurePosixPath(p).name, b.decode("utf-8")) for p, b in found]


def _figures(ctx: Context, figures_dir: str | None, manuscript: str) -> dict[str, str]:
    """Figure id (file-name stem up to the first '_') -> path relative to the manuscript's folder."""
    if not figures_dir:
        return {}
    by_id: dict[str, str] = {}
    for name in sorted(ctx.remote.listdir(figures_dir.rstrip("/")),
                       key=lambda n: next((i for i, ext in enumerate(IMAGE_PREFERENCE) if n.lower().endswith(ext)), 9)):
        if not name.lower().endswith(IMAGE_PREFERENCE):
            continue
        fid = PurePosixPath(name).stem.split("_")[0]
        by_id.setdefault(fid, f"{figures_dir.rstrip('/')}/{name}")
    base = PurePosixPath(manuscript).parent
    return {fid: _relpath(path, base) for fid, path in by_id.items()}


def _relpath(path: str, base: PurePosixPath) -> str:
    p, b = PurePosixPath(path).parts, base.parts
    common = 0
    while common < min(len(p), len(b)) and p[common] == b[common]:
        common += 1
    return "/".join([".."] * (len(b) - common) + list(p[common:]))


def build(ctx: Context, m) -> dict:
    """Assemble the English manuscript; returns {text, markers, missing, warnings, used, inputs}."""
    tpl, out_path = m.paths.manuscript_templates, m.paths.manuscript_out.replace("{lang}", "en")
    a = m.assembly
    wanted = [tpl] + [p for p in (m.paths.own_evidence, a.references_csv, a.table_specs, m.paths.bib,
                                  m.paths.captions, m.paths.tables) if p] + list(a.sections.values())
    files = ctx.remote.pull(sorted({w.rstrip("/") for w in wanted}))
    tables = _under(files, m.paths.tables) if m.paths.tables else {}

    def table_csv(tid: str) -> str | None:
        key = f"{m.paths.tables.rstrip('/')}/{tid}.csv" if m.paths.tables else None
        return tables[key].decode("utf-8") if key in tables else None

    inputs_sha = {p: sha256(b) for p, b in files.items() if not p.startswith((m.paths.tables or "\0"))}
    templates = _templates(files, tpl.rstrip("/"))
    if not templates:
        raise AssemblyError(f"no templates at {tpl}")
    if m.placeholders.dialect == "floodstate_fill":
        text, missing, n = F.fill(templates[0][1], table_csv)
        return {"text": text, "markers": markers.find_marker_ids(text), "missing": missing, "warnings": [],
                "used": [], "placeholders": n, "inputs": inputs_sha, "out_path": out_path}

    def frame(path: str | None, index: str) -> pd.DataFrame | None:
        if not path or path not in files:
            return None
        return pd.read_csv(io.BytesIO(files[path]), dtype=str, keep_default_na=False).set_index(index)

    bib = None
    if not a.references_csv and m.paths.bib in files:
        from src.services.bibtex import parse_bib
        bib = {e["key"]: e for e in parse_bib(files[m.paths.bib].decode("utf-8"))}
    captions = parse_captions(files[m.paths.captions].decode("utf-8")) if m.paths.captions in files else {}
    inputs = Inputs(
        evidence=frame(m.paths.own_evidence, "claim_id"), references=frame(a.references_csv, "cite_key"), bib=bib,
        sections={name: files[p].decode("utf-8") for name, p in a.sections.items() if p in files},
        table_specs=yaml.safe_load(files[a.table_specs]) if a.table_specs in files else {},
        table_csv=table_csv, captions=captions, figures=_figures(ctx, m.paths.figures, out_path))
    asm = Assembler(inputs, number_tables=a.number_tables, number_figures=a.number_figures,
                    reference_source=PurePosixPath(a.references_csv).name if a.references_csv else "the .bib")
    body = asm.assemble(templates)
    if inputs.references is not None:
        refs = asm.references_csv_section()
    else:
        from src.workbench import client
        rendered = client.api().bib.render(files[m.paths.bib].decode("utf-8"), keys=sorted(asm.cited),
                                           style=m.citation.style) if asm.cited else {"references": []}
        refs = "# REFERENCES\n\n" + "\n".join(f"- {r['text']}" for r in rendered["references"]) + "\n"
    text = body + "\n" + refs
    return {"text": text, "markers": markers.find_marker_ids(text), "missing": [], "warnings": asm.warnings,
            "used": sorted(asm.used), "numbers": {f"{k}:{i}": n for (k, i), n in asm.numbers.items()},
            "inputs": inputs_sha, "out_path": out_path}


def run(ctx: Context, *, final: bool = False) -> int:
    m = ctx.manifest()
    if m is None:
        print(f"{ctx.project_id}: no manifest; run init first")
        return 1
    if m.placeholders.dialect == "none" or not m.paths.manuscript_templates or not m.paths.manuscript_out:
        print(f"{ctx.project_id}: the manuscript is written directly (dialect {m.placeholders.dialect}); nothing to assemble")
        return 0
    try:
        result = build(ctx, m)
    except AssemblyError as exc:
        print(f"{ctx.project_id}: assembly failed: {exc}")
        return 1
    open_n, missing_n = len(result["markers"]), len(result["missing"])
    for w in result["warnings"]:
        print(f"  warning: {w}")
    if final and (open_n or missing_n):
        print(f"{ctx.project_id}: --final refused: {open_n} [PENDING] marker(s), {missing_n} [[MISSING]] cell(s); "
              "nothing staged")
        return 3
    out_path = result["out_path"]
    folder = str(PurePosixPath(out_path).parent)
    stage(ctx, out_path, result["text"])
    stage(ctx, f"{folder}/OPEN_ITEMS.md", open_items(result["text"]) +
          ("".join(f"- {x}\n" for x in result["missing"]) if result["missing"] else ""))
    if result["used"]:
        stage(ctx, f"{folder}/claims_used.csv", "claim_id\n" + "".join(f"{c}\n" for c in result["used"]))
    report = {k: v for k, v in result.items() if k != "text"} | {"final": final,
                                                                   "provenance": provenance(ctx, result["inputs"])}
    stage(ctx, f"{folder}/assembly_report.json", json.dumps(report, indent=1, ensure_ascii=False) + "\n")
    docx_path = m.assembly.docx or f"{folder}/private/{PurePosixPath(out_path).stem}.docx"
    stage(ctx, docx_path, to_docx(result["text"]))
    print(f"{ctx.project_id}: assembled {out_path} — {open_n} open marker(s), {missing_n} missing cell(s), "
          f"{len(result['used'])} claims; staged with OPEN_ITEMS.md and {docx_path}")
    return 0


def to_docx(markdown: str) -> bytes:
    from docx import Document

    from src.services.markdown_docx import convert_markdown
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "manuscript.md"
        md.write_text(markdown, encoding="utf-8")
        doc = Document()
        convert_markdown(doc, md)
        out = Path(tmp) / "manuscript.docx"
        doc.save(out)
        return out.read_bytes()
