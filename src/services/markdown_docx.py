"""Markdown → Word (.docx) rendering shared by the manuscript builders.

Moved out of the git-ignored ``src/paper_audit/article_1_final/build_docx.py`` so that
tracked code (``src/paper_3/v2/assemble.py``) does not depend on an ignored module.
Behaviour is unchanged: headings, paragraphs with inline **bold** / *italic*,
GitHub-style tables as native Word tables, and ``![alt](path)`` figures resolved
relative to the Markdown file. Unlike the original script it raises ImportError
instead of exiting the process when python-docx is missing.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

_INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
_IMG_RE = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.+?)\)\s*$")
_TEXT_WIDTH_IN = 6.3  # usable width on a portrait A4/Letter page with default margins


def add_runs(paragraph, text: str) -> None:
    """Append text to *paragraph*, honouring **bold** and *italic* spans."""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _cells(row: str) -> list[str]:
    row = row.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [c.strip() for c in row.split("|")]


def _is_separator(cells: list[str]) -> bool:
    return all(re.fullmatch(r":?-{1,}:?", c or "-") for c in cells)


def add_table(doc, rows: list[list[str]], font_size_pt: float | None = None) -> None:
    """Render cell rows (first row = header) as a Word table."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for ci in range(ncols):
            text = row[ci] if ci < len(row) else ""
            para = cells[ci].paragraphs[0]
            add_runs(para, text)
            for run in para.runs:
                if ri == 0:  # header row → bold
                    run.bold = True
                if font_size_pt is not None:
                    run.font.size = Pt(font_size_pt)
    doc.add_paragraph()  # spacer after table


def add_image(doc, path: Path, alt: str) -> None:
    if not path.exists():
        doc.add_paragraph(f"[missing image: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(_TEXT_WIDTH_IN))


def heading_level(hashes: int) -> int:
    """'#' → document title (level 0), '##' → level 1, '###' → level 2 …"""
    return max(0, hashes - 1)


def convert_markdown(doc, md_path: Path) -> None:
    """Walk *md_path* and append its content to the python-docx *doc*."""
    lines = md_path.read_text(encoding="utf-8").splitlines()
    base = md_path.parent
    i = 0
    n = len(lines)
    while i < n:
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).strip(), level=heading_level(len(m.group(1))))
            i += 1
            continue

        im = _IMG_RE.match(stripped)
        if im:
            add_image(doc, (base / im.group("path")).resolve(), im.group("alt"))
            i += 1
            continue

        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [r for r in (_cells(r) for r in block) if not _is_separator(r)]
            add_table(doc, rows)
            continue

        # ordinary paragraph — justified body text, 1.5 line spacing
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.line_spacing = 1.5
        add_runs(para, stripped)
        i += 1
